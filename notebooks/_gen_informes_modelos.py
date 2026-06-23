# -*- coding: utf-8 -*-
# Genera los 3 informes de modelos (CXR/ECG/LABS) actualizados a la estructura v1/v2/v3 con resultados y figuras.
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NBO=Path(r"C:\TFM\1.Opción - Symile Mimic\tfm_multimodal_clinico\notebooks")
DOCDIR=Path(r"C:\TFM\1.Opción - Symile Mimic\DOCUMENTACIÓN")
BLUE=RGBColor(0x1F,0x4E,0x79); LBLUE=RGBColor(0x2E,0x75,0xB6); GREEN=RGBColor(0x2E,0x5A,0x1E); GRAY=RGBColor(0x59,0x59,0x59); REDC=RGBColor(0xB0,0x2A,0x2A)
LAB=["Atelectasis","Cardiomegaly","Edema","Lung Opacity","No Finding","Pleural Effusion"]
ESL={"Atelectasis":"Atelectasia","Cardiomegaly":"Cardiomegalia","Edema":"Edema","Lung Opacity":"Opacidad pulmonar","No Finding":"No Finding","Pleural Effusion":"Derrame pleural"}

def newdoc():
    d=Document(); d.styles["Normal"].font.name="Calibri"; d.styles["Normal"].font.size=Pt(10.5); return d
def shade(c,h): sh=OxmlElement("w:shd"); sh.set(qn("w:val"),"clear"); sh.set(qn("w:fill"),h); c._tc.get_or_add_tcPr().append(sh)
def H1(d,t):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(15); p.paragraph_format.space_after=Pt(5); r=p.add_run(t); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=BLUE
def H2(d,t):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(9); p.paragraph_format.space_after=Pt(3); r=p.add_run(t); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=LBLUE
def P(d,t,sp=4,it=False,color=None,bold=False):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(sp); r=p.add_run(t); r.italic=it; r.bold=bold
    if color: r.font.color.rgb=color
def bullet(d,t):
    p=d.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(2)
    for i,seg in enumerate(t.split("**")):
        r=p.add_run(seg); r.bold=(i%2==1)
def img(d,path,w=15.0,cap=None):
    path=NBO/path
    if not path.exists(): P(d,f"[figura no disponible: {path.name}]",it=True,color=REDC); return
    d.add_picture(str(path),width=Cm(w)); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    if cap:
        c=d.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=c.add_run(cap); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=GRAY
def table(d,headers,rows,best=None,small=False):
    t=d.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; fs=Pt(8.5) if small else Pt(9.5)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,"1F4E79"); pr=c.paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=pr.add_run(h); run.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=fs
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            pr=cells[i].paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.CENTER if i>0 else WD_ALIGN_PARAGRAPH.LEFT; run=pr.add_run(str(v)); run.font.size=fs
            if best is not None and ri==best: run.bold=True; shade(cells[i],"E2EFDA")
    d.add_paragraph().paragraph_format.space_after=Pt(2)

def decisiones_comunes(d, extra):
    H1(d,"2. Decisiones críticas (comunes a todo el TFM)")
    H2(d,"2.1 Problema multietiqueta con enmascarado (masking)")
    P(d,"Cada paciente puede tener varios hallazgos → 6 salidas independientes con sigmoide y pérdida BCE enmascarada.")
    bullet(d,"**NaN** (no anotado) → se **enmascara**: no entra en la pérdida ni en la métrica de esa etiqueta.")
    bullet(d,"**−1** (incertidumbre CheXpert) → política `uncertainty_policy` (por defecto 0).")
    H2(d,"2.2 'No Finding' y negativos derivados por exclusividad (innovación clave)")
    P(d,"Atelectasia, Lung Opacity y No Finding casi no tenían negativos anotados → no entrenables. Se derivan negativos por exclusividad clínica, rellenando SOLO NaN:")
    bullet(d,"Si **No Finding = 1** → patologías con NaN pasan a **0 observado**.")
    bullet(d,"Si **alguna patología = 1** → No Finding con NaN pasa a **0 observado**.")
    bullet(d,"Nunca se pisa un valor explícito ni un −1. Esto las convierte en entrenables sin inventar etiquetas.")
    H2(d,"2.3 Desequilibrio, calibración y OOF")
    bullet(d,"**Desequilibrio**: `pos_weight`/`scale_pos_weight` por etiqueta; umbral por F1 en validación.")
    bullet(d,"**Calibración isotónica** (ajustada en val): las probabilidades se leen como tales.")
    bullet(d,"**OOF sin fuga**: cada muestra de train la predice un modelo que no la vio → features válidas para la fusión.")
    if extra:
        H2(d,extra[0]);
        for b in extra[1]: bullet(d,b)

def build(cfg):
    d=newdoc()
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(cfg["titulo"]); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=BLUE
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(cfg["subtitulo"]); r.font.size=Pt(11); r.font.color.rgb=LBLUE
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("TFM · Universidad de Salamanca · Documento actualizado (estructura v1/v2/v3)"); r.font.size=Pt(9); r.font.color.rgb=GRAY
    H1(d,"1. Rol de la modalidad y estrategia")
    for b in cfg["rol"]: bullet(d,b)
    decisiones_comunes(d, cfg.get("extra"))
    H1(d,"3. Las 3 versiones (ligera / óptima / pesada)")
    table(d,["Versión","Estructura","Test macro AUC"],cfg["versiones"],best=cfg["best_idx"])
    P(d,cfg["nota_versiones"],it=True,color=GRAY)
    H1(d,"4. Modelo elegido y justificación")
    P(d,cfg["justif"], bold=False)
    H2(d,"Resultados del modelo elegido (AUC-ROC por etiqueta, test)")
    table(d,["Hallazgo","AUC"],[[ESL[l],("%.3f"%cfg["perlabel"][l] if cfg["perlabel"][l] is not None else "N/A")] for l in LAB])
    P(d,f"MACRO patologías = {cfg['macro']}", bold=True, color=GREEN)
    H1(d,"5. Figuras")
    for fpath,cap,w in cfg["figs"]: img(d,fpath,w,cap)
    H1(d,"6. Conclusión")
    P(d,cfg["conclusion"])
    P(d,"")
    P(d,"Mantiene todas las bases: masking, negativos derivados, calibración isotónica y export de OOF para el stacking.",it=True,color=GRAY)
    out=DOCDIR/cfg["fichero"]
    try: d.save(str(out)); print("Guardado:",cfg["fichero"])
    except PermissionError:
        alt=out.with_name(out.stem+" (nuevo).docx"); d.save(str(alt)); print("BLOQUEADO ->",alt.name)

CXR=dict(
 titulo="Informe del Modelo de Imagen — Radiografía de tórax (CXR)",
 subtitulo="DenseNet121 congelado + embeddings cacheados · 3 versiones",
 fichero="Documentación Informe Modelo CXR (actualizado v1-v3).docx",
 rol=["**Modalidad principal**: las 6 etiquetas son hallazgos radiográficos (CheXpert) → la imagen es casi su propio oráculo.",
      "**Estrategia CPU**: backbone **DenseNet121** (`densenet121-res224-all`, CheXNet) **congelado**; se pasan las imágenes una vez, se cachean los **embeddings (1024 dims)** y se entrena una cabeza encima."],
 extra=("2.4 El fallo del CXR (AUC≈0.5) y su corrección — específico de imagen",
        ["Causa: preprocesado mal alineado (z-score ImageNet, 3→1 canal, 160×160) → entradas que el backbone nunca vio.",
         "Corrección: **min-max por imagen a [−1024, 1024]**, 1 canal, **224×224** (lo que espera torchxrayvision). El AUC pasó de 0.5 a 0.81.",
         "El fine-tuning agresivo (versiones previas) **empeoraba** → se mantiene congelado."]),
 versiones=[["v1 ligera","LogReg one-vs-rest sobre embeddings","0.731"],
            ["v2 ÓPTIMA","Cabeza MLP + metadatos + correlación de etiquetas + K-fold","0.809"],
            ["v3 pesada","DenseNet+TTA(flip)+ResNet50@512 + bagging + ensemble LogReg","(pendiente)"]],
 best_idx=1, macro="0.809",
 nota_versiones="El salto v1→v2 (LogReg→MLP+metadatos) es el grande; la v3 explora diversidad de backbones pero el techo de estas etiquetas ya lo marca el DenseNet congelado.",
 justif="Se elige la **v2 (óptima)**: mejor AUC macro (0.809), buena calibración y coste contenido (la pasada del backbone se cachea; la cabeza se entrena en minutos). Bate claramente a la v1 ligera (0.731) y a las versiones pesadas con fine-tuning que se probaron y empeoraban.",
 perlabel={"Atelectasis":0.782,"Cardiomegaly":0.776,"Edema":0.832,"Lung Opacity":0.797,"No Finding":0.764,"Pleural Effusion":0.857},
 figs=[("outputs_cxr_densenet121_v2/results_v5.png","v2 óptima — resumen de resultados (AUC por etiqueta, test).",15.0),
       ("outputs_cxr_densenet121_v1/figuras/auc_por_etiqueta_v1.png","v1 ligera — AUC por etiqueta (LogReg sobre embeddings).",13.0),
       ("outputs_cxr_densenet121_v1/figuras/confusion_v1.png","v1 ligera — matrices de confusión (test).",15.0)],
 conclusion="El módulo de imagen es el más fuerte (0.809) y sostiene el sistema. La clave no fue la complejidad sino el preprocesado correcto + DenseNet121 congelado + los negativos derivados de No Finding.")

ECG=dict(
 titulo="Informe del Modelo de Señal — Electrocardiograma (ECG)",
 subtitulo="ResNet1D desde cero, 12 derivaciones · 3 versiones",
 fichero="Documentación Informe Modelo ECG (actualizado v1-v3).docx",
 rol=["**Modalidad complementaria**: el ECG da señal **indirecta** sobre hallazgos radiográficos (techo ~0.64 AUC).",
      "**Arquitectura**: ResNet1D entrenada desde cero, 12 derivaciones como canales de entrada, z-score por derivación, aumento de señal (ruido, escalado, desplazamiento, lead-masking)."],
 extra=None,
 versiones=[["v1 ligera","ResNet1D base32, 1 bloque/etapa, señal 1000","(pendiente)"],
            ["v2 ÓPTIMA","base48, 2 bloques/etapa, señal 1250, Optuna, K-fold ensemble","0.637"],
            ["v3 pesada","base64, señal 2500, K=5, aumento rico","0.639"]],
 best_idx=1, macro="0.637",
 nota_versiones="v2 y v3 quedan prácticamente empatadas (0.637 vs 0.639) pese a que la v3 cuesta muchas más horas → la señal del ECG tiene techo.",
 justif="Se elige la **v2 (óptima)** como módulo de referencia: prácticamente igual que la v3 pesada (0.637 vs 0.639) a una fracción del coste. Es la que alimenta el stacking. La v3 se conserva como variante de máximo nivel pero su ganancia es marginal.",
 perlabel={"Atelectasis":0.633,"Cardiomegaly":0.676,"Edema":0.628,"Lung Opacity":0.635,"No Finding":0.626,"Pleural Effusion":0.615},
 figs=[("outputs_ecg_resnet1d_v2/figuras/auc_por_etiqueta_v3_1.png","v2 óptima — AUC por etiqueta (test).",13.0),
       ("outputs_ecg_resnet1d_v2/figuras/curvas_v3_1.png","v2 óptima — curvas ROC/PR/calibración y entrenamiento.",15.0),
       ("outputs_ecg_resnet1d_v2/figuras/confusion_matrices_v3_1.png","v2 óptima — matrices de confusión (test).",15.0)],
 conclusion="El ECG aporta poco en solitario, pero es valioso como modalidad complementaria: en la herramienta de decisión es la prueba que más 'manda' en Cardiomegalia (~15% de los pacientes).")

LABS=dict(
 titulo="Informe del Modelo Tabular — Analíticas de sangre",
 subtitulo="Ensemble multiarquitectura · 3 versiones",
 fichero="Documentación Informe Modelo LABS Tabular (actualizado v1-v3).docx",
 rol=["**Modalidad complementaria**: techo de la señal tabular ~0.69 AUC.",
      "**Diseño**: árboles (XGBoost/LightGBM) con **NaN nativo** sobre el valor bruto; modelos densos (LogReg/MLP) sobre el percentil imputado **por fold**."],
 extra=("2.4 Decisiones específicas del tabular",
        ["**Ratios clínicos** con sentido fisiológico: BUN/Creatinina, Neutrófilos/Linfocitos, RDW×Edad.",
         "**Flags de missingness** (que falte una analítica es informativo → MNAR).",
         "**Sin `cxr_view`** (fuga de la modalidad imagen) e **imputación por fold** (sin leakage)."]),
 versiones=[["v1 ligera","Un XGBoost one-vs-rest (NaN nativo) + ratios + flags","0.673"],
            ["v2 ÓPTIMA","Ensemble XGBoost+LightGBM+LogReg+MLP, tuning por modelo","0.692"],
            ["v3 pesada","Ensemble + TabPFN (nube) + stacking de 2º nivel","(pendiente)"]],
 best_idx=1, macro="0.692",
 nota_versiones="El ensemble (v2) aporta +0.02 sobre el XGBoost único (v1); el techo de la señal tabular limita ganancias mayores.",
 justif="Se elige la **v2 (óptima)**: el **ensemble por promedio** de 4 arquitecturas tuneadas (0.692) bate al XGBoost único (0.673) y es el que alimenta el stacking. La v3 (TabPFN + stacking de 2º nivel) se conserva como máximo nivel.",
 perlabel={"Atelectasis":0.695,"Cardiomegaly":0.623,"Edema":0.735,"Lung Opacity":0.702,"No Finding":0.701,"Pleural Effusion":0.704},
 figs=[("outputs_labs_tabular_v2/figuras/boxplot_auc_modelos.png","v2 óptima — comparativa de arquitecturas + ensemble (CV).",14.0),
       ("outputs_labs_tabular_v2/figuras/importancia_variables.png","v2 óptima — importancia de variables (XGBoost, nombres legibles).",13.0),
       ("outputs_labs_tabular_v2/figuras/confusion_mejor.png","v2 óptima — matrices de confusión del mejor (test).",15.0)],
 conclusion="Las analíticas son la mejor modalidad complementaria: en la herramienta de decisión lideran en Edema (~16%) y Derrame, donde el líquido deja huella en sangre.")

for cfg in (CXR,ECG,LABS): build(cfg)
print("=== 3 informes de modelos generados ===")
