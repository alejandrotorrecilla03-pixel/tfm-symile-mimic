# -*- coding: utf-8 -*-
# Dosier de USO de la herramienta de decisión: explica cada número/tabla/línea y recorre 6 casos.
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
NBO=Path(r"C:\TFM\1.Opción - Symile Mimic\tfm_multimodal_clinico\notebooks")
OUT=NBO/"outputs_herramienta_decision"; DOS=OUT/"dosier"; FGE=OUT/"figuras"
DOC=Path(r"C:\TFM\1.Opción - Symile Mimic\DOCUMENTACIÓN\Dosier de Uso - Herramienta de Decisión Clínica.docx")
D=json.load(open(OUT/"dosier_casos.json",encoding="utf-8"))
MODS=["CXR","ECG","LABS"]; MOD_ES={"CXR":"Radiografía","ECG":"ECG","LABS":"Analíticas"}
BLUE=RGBColor(0x1F,0x4E,0x79); LBLUE=RGBColor(0x2E,0x75,0xB6); GREEN=RGBColor(0x2E,0x5A,0x1E); GRAY=RGBColor(0x59,0x59,0x59); REDC=RGBColor(0xB0,0x2A,0x2A)
d=Document(); d.styles["Normal"].font.name="Calibri"; d.styles["Normal"].font.size=Pt(10.5)
def shade(c,h): sh=OxmlElement("w:shd"); sh.set(qn("w:val"),"clear"); sh.set(qn("w:fill"),h); c._tc.get_or_add_tcPr().append(sh)
def H1(t):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(15); p.paragraph_format.space_after=Pt(5); r=p.add_run(t); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=BLUE
def H2(t):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(9); p.paragraph_format.space_after=Pt(3); r=p.add_run(t); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=LBLUE
def P(t,sp=4,it=False,color=None,bold=False):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(sp)
    for i,seg in enumerate(t.split("**")):
        r=p.add_run(seg); r.bold=bold or (i%2==1); r.italic=it
        if color: r.font.color.rgb=color
def bullet(t):
    p=d.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(2)
    for i,seg in enumerate(t.split("**")): r=p.add_run(seg); r.bold=(i%2==1)
def img(path,w=15.5,cap=None):
    path=Path(path)
    if not path.exists(): P(f"[figura: {path.name}]",it=True,color=REDC); return
    d.add_picture(str(path),width=Cm(w)); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    if cap:
        c=d.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=c.add_run(cap); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=GRAY
def table(headers,rows,best=None,small=True,fills=None):
    t=d.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; fs=Pt(8.5) if small else Pt(9.5)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,"1F4E79"); pr=c.paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=pr.add_run(h); run.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=fs
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            pr=cells[i].paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.CENTER if i>0 else WD_ALIGN_PARAGRAPH.LEFT; run=pr.add_run(str(v)); run.font.size=fs
            if best is not None and ri==best: run.bold=True; shade(cells[i],"E2EFDA")
            if fills and (ri,i) in fills: shade(cells[i],fills[(ri,i)])
    d.add_paragraph().paragraph_format.space_after=Pt(2)
pct=lambda x: f"{round(100*x):d}%"
def pp(x):
    v=round(100*x); return ("+" if v>0 else "")+f"{v} pp"

def detalle(f):
    rows=[]
    for mo in MODS:
        rows.append([MOD_ES[mo], pct(f["probs"][mo]), f"{f['conf'][mo]:.2f}", f"{f['skill'][mo]:.2f}",
                     pct(f["w"][mo]), f"{f['aport'][mo]:.2f}", pct(f["contrafactual"][mo]), pp(f["influencia"][mo])])
    li=MODS.index(f["lead"])
    table(["Prueba","prob (p)","confianza","habilidad","peso","aportación","decisión sin ella","influencia"],rows,best=li)
    av=" · ⚠ baja certeza (ninguna prueba se moja)" if f["abstencion"] else (" · ⚠ discrepancia entre pruebas" if f["disagree"] else "")
    P(f"Decisión combinada = suma de aportaciones = **{pct(f['comb'])}** · prueba líder = **{MOD_ES[f['lead']]}**" +
      (f" · verdad clínica: **{'SÍ' if f['y']==1 else 'no'} {f['hallazgo']}**" if f.get('y') is not None else "") + av, sp=3)

# ===================== PORTADA =====================
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("Dosier de Uso — Herramienta de Apoyo a la Decisión Clínica"); r.bold=True; r.font.size=Pt(19); r.font.color.rgb=BLUE
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("Cómo leer cada número, tabla y línea · 6 casos de uso explicados"); r.font.size=Pt(12); r.font.color.rgb=LBLUE
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("TFM · CXR + ECG + Analíticas (modelos óptimos v2) · Universidad de Salamanca"); r.font.size=Pt(10); r.font.color.rgb=GRAY

# ===================== 1. QUÉ ES =====================
H1("1. Qué hace la herramienta")
P("Para un paciente concreto y cada hallazgo, la herramienta combina lo que dicen las tres pruebas (radiografía, ECG, "
  "analíticas) y **te dice en qué prueba fiarte más** y por qué, con una recomendación en lenguaje natural. No es un "
  "diagnóstico: es **apoyo** que surfacea y justifica la información.")

# ===================== 2. LOS 3 INGREDIENTES =====================
H1("2. Los tres ingredientes y la fórmula del peso")
table(["Ingrediente","Qué mide","Cómo se calcula","Rango"],
      [["Probabilidad (p)","qué dice cada prueba","salida calibrada del modelo (isotónica)","0 (no) … 1 (sí)"],
       ["Confianza","cuán segura está en ESTE paciente","|p − 0.5| · 2","0 (duda) … 1 (segura)"],
       ["Habilidad","en qué prueba fiarse para ESE hallazgo","AUC(validación) − 0.5","0 (azar) … 0.5 (perfecta)"]],small=False)
P("**Peso(prueba) = habilidad × confianza** (luego se normaliza para que los tres pesos sumen 100%). La prueba con más "
  "peso es la **líder** (la que mirar). La **decisión combinada** es la media de las probabilidades ponderada por el peso.")

# ===================== 3. CÓMO LEER (LEYENDA) =====================
H1("3. Cómo leer la figura (cada elemento)")
img(DOS/"leyenda_anotada.png",16.5,"Figura guía: cada elemento numerado se explica en la tabla siguiente.")
table(["Elemento","Qué es"],
      [["① Longitud de la barra","La probabilidad (p) que da esa prueba al hallazgo (0 a 1). Barra larga = la prueba cree que SÍ."],
       ["② 'peso X%'","Cuánto se le hace caso a esa prueba (= habilidad × confianza, normalizado). Es distinto de la barra."],
       ["③ Línea gris discontinua","El umbral 50%: a la izquierda 'probablemente no', a la derecha 'probablemente sí'."],
       ["④ Línea roja","La DECISIÓN combinada (media ponderada de las barras por su peso)."],
       ["⑤ Barra con borde negro","La prueba LÍDER: la de mayor peso, en la que conviene fijarse."]],small=False)

# ===================== 4. EXPLICABILIDAD AVANZADA =====================
H1("4. Explicabilidad avanzada (cómo se justifica la decisión)")
P("Además de los tres ingredientes, en cada caso se muestra una tabla con columnas adicionales que **descomponen** la decisión:")
bullet("**Aportación**: cuánto pone cada prueba en la decisión final (= peso × probabilidad). Las tres aportaciones **suman exactamente la decisión**.")
bullet("**Decisión sin ella** (contrafactual): qué saldría si **quitáramos** esa prueba y solo usáramos las otras dos.")
bullet("**Influencia**: cuánto cambia la decisión esa prueba (= decisión − decisión sin ella), en puntos porcentuales (pp). Mide su peso real en el resultado.")
bullet("**Avisos**: ⚠ **discrepancia** (pruebas con habilidad que se contradicen) y ⚠ **baja certeza** (ninguna prueba se moja → tratar con cautela).")
P("La fiabilidad de la prueba líder se estima además con su acierto histórico a esa confianza (validación): a más "
  "confianza, más acierto, como muestra esta curva.")
img(FGE/"exp_acierto_vs_confianza.png",10.5,"A mayor confianza de una prueba, mayor acierto (validación).")

# ===================== 5. CASOS DE USO =====================
H1("5. Casos de uso (paso a paso)")
NOM={"A":"Caso A — La radiografía es clara: guíate por ella",
     "B":"Caso B — La radiografía duda y otra prueba acierta",
     "C":"Caso C — Las pruebas se contradicen: manda la más fiable",
     "D":"Caso D — Una prueba muy segura pero poco hábil: se le baja el peso",
     "E":"Caso E — Baja certeza global: ninguna prueba se moja (precaución)",
     "F":"Caso F — Paciente con varios hallazgos a la vez"}
LECCION={"A":"Cuando la radiografía está segura y es la más hábil, lidera y la decisión la sigue. Caso típico y mayoritario.",
     "B":"Si la radiografía no se moja (confianza baja), su peso se desploma y manda la prueba más segura y hábil para ese hallazgo (a menudo las analíticas en Edema/Derrame).",
     "C":"Cuando las pruebas se contradicen, la herramienta avisa y se inclina por la de mayor habilidad. Es donde más ayuda al clínico a no dejarse llevar por una prueba ruidosa.",
     "D":"Una prueba puede estar muy segura y aun así pesar poco si su habilidad para ese hallazgo es baja (típico del ECG). Confianza no es lo mismo que acierto.",
     "E":"Si ninguna prueba se moja, la herramienta marca baja certeza: la decisión es poco fiable y conviene apoyarse en la clínica.",
     "F":"La herramienta razona hallazgo por hallazgo; un mismo paciente puede tener una prueba líder distinta en cada uno."}
casos={c["_caso"]:c for c in D["casos"]}
for code in ["A","B","C","D","E","F"]:
    if code not in casos: continue
    c=casos[code]; H2(NOM[code])
    P(c.get("_desc",""), it=True, color=GRAY)
    img(DOS/f"caso_{code}.png",15.5)
    if code=="F":
        for f in c["multi"]:
            P(f"Hallazgo: **{f['hallazgo']}**", sp=2); detalle(f)
    else:
        detalle(c)
    P("Lección: "+LECCION[code], color=GREEN, sp=8)

# ===================== 6. CÓMO EJECUTARLA =====================
H1("6. Cómo ejecutarla")
bullet("**App gráfica** (recomendada): doble clic en `Lanzar_Herramienta_Clinica.bat` → se abre en el navegador (http://localhost:8501). Eliges un paciente y ves estas figuras y tablas en vivo.")
bullet("**API** (uso programático): `uvicorn api_decision:app --host 127.0.0.1 --port 8000` → documentación en /docs; `GET /paciente/{hadm_id}` devuelve todo en JSON.")
bullet("Ambas usan las salidas de los **modelos óptimos v2** (CXR/ECG/analíticas). Datos derivados de MIMIC → uso local (127.0.0.1).")

# ===================== 7. HONESTIDAD =====================
H1("7. Honestidad y límites")
bullet("Es **apoyo, no diagnóstico**: la decisión combinada acierta más que cualquier prueba sola, pero también falla (ver casos).")
bullet("Como las etiquetas vienen del informe del CXR, la herramienta **suele recomendar la radiografía**; su valor está en los casos de CXR dudoso (Edema/Derrame → analíticas; Cardiomegalia → ECG).")
bullet("La habilidad y la fiabilidad se estiman en **validación** (n limitado): son orientación estadística, no garantías individuales.")

try: d.save(str(DOC)); print("Guardado:",DOC.name)
except PermissionError:
    alt=DOC.with_name(DOC.stem+" (nuevo).docx"); d.save(str(alt)); print("BLOQUEADO ->",alt.name)
print("imágenes:",len(d.inline_shapes),"· tablas:",len(d.tables))
