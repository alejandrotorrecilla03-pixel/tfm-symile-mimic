# -*- coding: utf-8 -*-
# Informe del stacking actualizado: base (v1) + avanzado (v2), con resultados reales y figuras.
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
NBO=Path(r"C:\TFM\1.Opción - Symile Mimic\tfm_multimodal_clinico\notebooks")
DOC=Path(r"C:\TFM\1.Opción - Symile Mimic\DOCUMENTACIÓN\Documentación Informe Stacking Multimodal (actualizado).docx")
BLUE=RGBColor(0x1F,0x4E,0x79); LBLUE=RGBColor(0x2E,0x75,0xB6); GREEN=RGBColor(0x2E,0x5A,0x1E); GRAY=RGBColor(0x59,0x59,0x59); REDC=RGBColor(0xB0,0x2A,0x2A)
d=Document(); d.styles["Normal"].font.name="Calibri"; d.styles["Normal"].font.size=Pt(10.5)
def shade(c,h): sh=OxmlElement("w:shd"); sh.set(qn("w:val"),"clear"); sh.set(qn("w:fill"),h); c._tc.get_or_add_tcPr().append(sh)
def H1(t):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(15); p.paragraph_format.space_after=Pt(5); r=p.add_run(t); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=BLUE
def H2(t):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(9); p.paragraph_format.space_after=Pt(3); r=p.add_run(t); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=LBLUE
def P(t,sp=4,it=False,color=None,bold=False):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(sp); r=p.add_run(t); r.italic=it; r.bold=bold
    if color: r.font.color.rgb=color
def bullet(t):
    p=d.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(2)
    for i,seg in enumerate(t.split("**")): r=p.add_run(seg); r.bold=(i%2==1)
def img(path,w=15.0,cap=None):
    path=NBO/path
    if not path.exists(): P(f"[figura no disponible: {path.name}]",it=True,color=REDC); return
    d.add_picture(str(path),width=Cm(w)); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    if cap:
        c=d.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=c.add_run(cap); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=GRAY
def table(headers,rows,best=None,small=False):
    t=d.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; fs=Pt(8.5) if small else Pt(9.5)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,"1F4E79"); pr=c.paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=pr.add_run(h); run.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=fs
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            pr=cells[i].paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.CENTER if i>0 else WD_ALIGN_PARAGRAPH.LEFT; run=pr.add_run(str(v)); run.font.size=fs
            if best is not None and ri==best: run.bold=True; shade(cells[i],"E2EFDA")
    d.add_paragraph().paragraph_format.space_after=Pt(2)

p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("Informe de Fusión Multimodal — Stacking (CXR + ECG + Analíticas)"); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=BLUE
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("Dos niveles: stacking BASE (v1) y stacking AVANZADO (v2) · Documento actualizado"); r.font.size=Pt(11); r.font.color.rgb=LBLUE

H1("1. Qué fusiona y con qué entradas")
bullet("Combina las predicciones de los **tres modelos óptimos (v2)**: CXR (0.809), ECG (0.637) y LABS (0.692).")
bullet("Entradas: probabilidades **calibradas** y **out-of-fold (OOF, sin fuga)** de cada modalidad (`*_pred_*.csv`).")
bullet("Sin fuga: para train se usan OOF; para test los modelos base se reentrenan en todo el train y predicen un test limpio.")

H1("2. Stacking BASE (v1): meta-modelo riguroso + interpretación")
P("Meta-modelos Regresión Logística (L2) y XGBoost (poco profundo) sobre las 18 probabilidades + contexto clínico, "
  "comparados con el blending (promedio) y con el mejor modelo individual (CXR).")
table(["Enfoque","macro AUC (patologías, test)"],
      [["CXR en solitario","0.7825"],["Blending (media simple)","0.7738"],["Stacking Regresión Logística","0.7869"],["Stacking XGBoost","0.7875"]],best=3)
bullet("El **blending simple empeora** al CXR: lo diluye con ECG/LABS más débiles.")
bullet("El **stacking gana poco** (+0.005) → confirma que la fusión sobre estas etiquetas tiene techo.")
img("outputs_stacking_v1/figuras/comparativa_enfoques.png",15.0,"Base — comparativa de enfoques y AUC por etiqueta (test).")
H2("Interpretación: ¿qué modalidad pesa en cada enfermedad?")
P("El meta-modelo permite leer el peso de cada modalidad por hallazgo: CXR domina; las analíticas aportan en Edema/Derrame; el ECG en Cardiomegalia.")
img("outputs_stacking_v1/figuras/pesos_modalidad_enfermedad.png",11.0,"Base — peso de cada modalidad por enfermedad (coeficientes del meta-modelo).")

H1("3. Stacking AVANZADO (v2): exprimir el AUC")
P("Sobre las mismas entradas se aplican técnicas más sofisticadas, eligiendo en validación y reportando en test:")
bullet("**Blending óptimo ponderado** (Nelder-Mead por etiqueta) y su variante en **espacio logit** / con probas **raw**.")
bullet("**Stacking enriquecido** (probas calibradas + raw + confianza + entropía).")
bullet("**Mixture-of-Experts** (gating por paciente) · **selección por etiqueta** · **ensemble de Caruana** (hill-climbing).")
table(["Enfoque de fusión","macro AUC (test)"],
      [["CXR en solitario","0.7825"],["Blending óptimo ponderado","0.7914"],["Blending en logit","0.7936"],["Stacking enriquecido","0.7903"],
       ["Mixture-of-Experts","0.7935"],["Selección por etiqueta","0.7935"],["Ensemble de Caruana (mejor)","0.7946"]],best=6)
img("outputs_stacking_v2/figuras/comparativa_v4.png",13.0,"Avanzado — comparativa de técnicas (test).")
img("outputs_stacking_v2/figuras/auc_etiqueta_v4.png",15.0,"Avanzado — AUC por etiqueta y técnica.")

H1("4. Por qué la fusión sube poco (y dónde sí aporta)")
bullet("Las 6 etiquetas se extraen del **informe del CXR** → la imagen es casi su propio oráculo; ECG y analíticas solo añaden señal indirecta.")
bullet("La mejora total es **modesta** (de 0.783 a 0.795, +0.012) pero **real** y rigurosa (OOF, test limpio, comparación con baselines).")
bullet("**Dónde aporta**: Edema y Derrame (analíticas) y Cardiomegalia (ECG). Esta lectura es la base de la herramienta de decisión clínica.")

H1("5. Conclusión")
P("El stacking aporta una mejora pequeña pero metodológicamente sólida; su mayor valor es **interpretativo**: muestra qué "
  "prueba pesa en cada enfermedad, lo que se convierte en el producto clínico del TFM (la herramienta de apoyo a la decisión).")
P("")
P("Mejor enfoque global: **ensemble de Caruana (0.7946)**. Entradas: modelos óptimos v2; selección en validación; test limpio.",it=True,color=GRAY)

try: d.save(str(DOC)); print("Guardado:",DOC.name)
except PermissionError:
    alt=DOC.with_name(DOC.stem+" (nuevo).docx"); d.save(str(alt)); print("BLOQUEADO ->",alt.name)
print("imgs:",len(d.inline_shapes),"tablas:",len(d.tables))
