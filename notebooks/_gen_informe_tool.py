# -*- coding: utf-8 -*-
# Informe didáctico de la herramienta de decisión: explica calibración, confianza y habilidad con ejemplos.
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NBO=Path(r"C:\TFM\1.Opción - Symile Mimic\tfm_multimodal_clinico\notebooks")
FG=NBO/"outputs_herramienta_decision"/"figuras"
DOC=Path(r"C:\TFM\1.Opción - Symile Mimic\DOCUMENTACIÓN\Documentación Informe Herramienta Decisión Clínica.docx")
D=json.load(open(NBO/"outputs_herramienta_decision"/"exploracion.json",encoding="utf-8"))
ES={"Atelectasis":"Atelectasia","Cardiomegaly":"Cardiomegalia","Edema":"Edema","Lung Opacity":"Opacidad pulmonar","No Finding":"Sin hallazgos","Pleural Effusion":"Derrame pleural"}
PATHO=["Atelectasis","Cardiomegaly","Edema","Lung Opacity","Pleural Effusion"]; MODS=["CXR","ECG","LABS"]; MOD_ES={"CXR":"Radiografía","ECG":"ECG","LABS":"Analíticas"}
BLUE=RGBColor(0x1F,0x4E,0x79); LBLUE=RGBColor(0x2E,0x75,0xB6); GREEN=RGBColor(0x37,0x56,0x23); GRAY=RGBColor(0x59,0x59,0x59); REDC=RGBColor(0xC0,0x00,0x00)
doc=Document(); doc.styles["Normal"].font.name="Calibri"; doc.styles["Normal"].font.size=Pt(10.5)
def shade(c,h): sh=OxmlElement("w:shd"); sh.set(qn("w:val"),"clear"); sh.set(qn("w:fill"),h); c._tc.get_or_add_tcPr().append(sh)
def H(t,size=15,color=BLUE,bef=14,aft=6):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(bef); p.paragraph_format.space_after=Pt(aft); r=p.add_run(t); r.bold=True; r.font.size=Pt(size); r.font.color.rgb=color; return p
def P(t,size=10.5,it=False,sp=4,color=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(sp); r=p.add_run(t); r.font.size=Pt(size); r.italic=it
    if color: r.font.color.rgb=color
    return p
def bullet(t):
    p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(2); p.add_run(t)
def img(path,w=15.5,cap=None):
    if not Path(path).exists(): P(f"[falta {Path(path).name}]",it=True); return
    doc.add_picture(str(path),width=Cm(w)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    if cap:
        c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=c.add_run(cap); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=GRAY
def table(headers,rows,best=None,small=False,fill_head="1F4E79"):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; fs=Pt(8.5) if small else Pt(9.5)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,fill_head); pr=c.paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=pr.add_run(h); run.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=fs
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            pr=cells[i].paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.CENTER if i>0 else WD_ALIGN_PARAGRAPH.LEFT; run=pr.add_run(str(v)); run.font.size=fs
            if best is not None and ri==best: run.bold=True; shade(cells[i],"E2EFDA")
    return t
pct=lambda x: f"{round(100*x):d}%"

# ===== PORTADA =====
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("Herramienta de Apoyo a la Decisión Clínica Multimodal"); r.bold=True; r.font.size=Pt(19); r.font.color.rgb=BLUE
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("¿En qué prueba fiarse para cada paciente y hallazgo? · Explicabilidad con ejemplos"); r.font.size=Pt(12); r.font.color.rgb=LBLUE
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("TFM · CXR + ECG + Analíticas · Universidad de Salamanca"); r.font.size=Pt(10); r.font.color.rgb=GRAY

# ===== 0. IDEA =====
H("1. La idea: del AUC a la decisión")
P("Las 6 etiquetas son hallazgos extraídos del informe de la radiografía, así que la radiografía es casi su propio "
  "oráculo y combinar modelos sube poco el AUC (de 0.783 a ~0.795). La conclusión útil no es el número global, sino "
  "responder, paciente a paciente: «para este hallazgo, ¿en qué prueba debe fijarse más el médico?» — y justificarlo.")
P("Para ello la herramienta combina, por hallazgo y paciente, tres ingredientes que ya tenemos. Este informe explica "
  "cada uno con ejemplos reales del conjunto de test.")
table(["Ingrediente","Qué mide","Fórmula"],
      [["1. Probabilidad calibrada","Qué dice cada prueba","p (ya calibrada por isotónica)"],
       ["2. Confianza","Cuán segura está esa prueba en ESTE paciente","|p − 0.5| · 2"],
       ["3. Habilidad","En qué prueba fiarse para ESE hallazgo","AUC(validación) − 0.5"]])

# ===== 2. PROBABILIDAD CALIBRADA =====
H("2. Probabilidad calibrada — «qué dice cada prueba»")
P("Cada modelo (radiografía, ECG, analíticas) produce, para cada hallazgo, un número entre 0 y 1. Pero la salida cruda "
  "de un modelo no es una probabilidad real. Por eso se calibra con regresión isotónica ajustada en validación: tras "
  "calibrar, un 0.80 significa que, entre casos donde la prueba dijo ~0.80, aproximadamente el 80% resultaron positivos. "
  "Eso permite LEER el número como una probabilidad clínica y compararlo entre pruebas.")
cx=dict(D["calib"]["CXR"])
near=min(D["calib"]["CXR"],key=lambda pair: abs(pair[0]-0.8))
P(f"Ejemplo (radiografía, agregando patologías en test): cuando dice ≈{near[0]:.2f}, la frecuencia real de positivos "
  f"observada fue ≈{near[1]:.2f}. La curva siguiente muestra que las tres pruebas están razonablemente calibradas "
  "(cerca de la diagonal); la radiografía es la mejor calibrada y el ECG la que más se desvía.")
img(FG/"exp_calibracion.png",11,"Figura 1. Calibración: probabilidad que dice la prueba (x) frente a frecuencia real de positivos (y).")

# ===== 3. CONFIANZA =====
H("3. Confianza — «cuán segura está en este paciente»")
P("La confianza es la distancia de la probabilidad al 0.5: conf = |p − 0.5| · 2. Un 0.95 (o un 0.05) es una prueba muy "
  "segura (conf ≈ 0.9); un 0.52 es prácticamente una moneda al aire (conf ≈ 0.04). Distingue al paciente concreto: la "
  "misma prueba puede estar segurísima en un paciente y totalmente dudosa en otro.")
P("Lo importante es que la confianza SIRVE: a más confianza, más acierto. La figura lo demuestra en validación — cuando "
  "una prueba está muy segura, acierta mucho más que cuando duda. Por eso la usamos para decidir cuánto pesa cada prueba.")
img(FG/"exp_acierto_vs_confianza.png",10.5,"Figura 2. A mayor confianza, mayor acierto (validación, patologías agregadas).")
img(FG/"exp_confianza_hist.png",16,"Figura 3. Distribución de la confianza por prueba: la radiografía suele estar más segura; el ECG duda más a menudo.")

# ===== 4. HABILIDAD =====
H("4. Habilidad — «en qué prueba fiarse para ese hallazgo»")
P("Confianza no es lo mismo que acierto: una prueba puede estar muy segura y equivocarse sistemáticamente para cierto "
  "hallazgo. La habilidad mide, por prueba y hallazgo, cuánto mejor que el azar ordena los casos: habilidad = AUC(en "
  "validación) − 0.5 (0 = azar; 0.5 = perfecto). Es un valor fijo por hallazgo, calculado sin mirar al paciente.")
sk=D["skill"]
table(["Hallazgo","Radiografía","ECG","Analíticas"],
      [[ES[l],f"{sk['CXR'][l]:.2f}",f"{sk['ECG'][l]:.2f}",f"{sk['LABS'][l]:.2f}"] for l in PATHO])
P("Se lee con claridad: la radiografía es la más hábil en todos los hallazgos; las analíticas son la complementaria "
  "(suben en Edema y Derrame, donde el líquido deja huella en sangre); el ECG es el más débil salvo en Cardiomegalia, "
  "donde sí aporta. Esto es lo que evita que una prueba muy segura pero poco hábil (p. ej., un ECG convencido) mande.")
img(FG/"exp_habilidad.png",14,"Figura 4. Habilidad por prueba y hallazgo (AUC−0.5 en validación).")

# ===== 5. CÓMO SE COMBINAN =====
H("5. Cómo se combinan: peso = habilidad × confianza")
# buscar finding ilustrativo: lead LABS con CXR dudoso
wk=None
for e in D["ejemplos"]:
    for f in e["findings"]:
        if f["lead"]=="LABS" and f["conf"]["CXR"]<0.1:
            wk=(e,f); break
    if wk: break
P("Para cada hallazgo y paciente, cada prueba recibe un peso = habilidad × confianza, y se normaliza. La prueba con más "
  "peso es «la que mirar»; la decisión combinada es la media ponderada de las probabilidades. Veámoslo con un paciente real:")
if wk:
    e,f=wk; l=ES[f["label"]]
    P(f"Paciente hadm={e['hadm']}, hallazgo {l}:", size=10.5)
    rows=[]
    for mod in MODS:
        rows.append([MOD_ES[mod],f"{f['probs'][mod]:.2f}",f"{f['conf'][mod]:.2f}",f"{f['skill'][mod]:.2f}",
                     f"{f['skill'][mod]:.2f}×{max(f['conf'][mod],0.05):.2f}",pct(f["w"][mod])])
    table(["Prueba","prob (p)","confianza","habilidad","habilidad×conf",f"peso final"],rows,best=[m for m in MODS].index(f["lead"]))
    P(f"La radiografía dice {pct(f['probs']['CXR'])} pero con confianza {f['conf']['CXR']:.2f} (no se moja) → su peso se "
      f"desploma. Las analíticas están muy seguras ({f['conf']['LABS']:.2f}) y tienen habilidad decente ({f['skill']['LABS']:.2f}) "
      f"→ se llevan el {pct(f['w']['LABS'])} del peso y mandan. Decisión combinada {pct(f['comb'])}. "
      f"Verdad clínica: {'SÍ' if f['y']==1 else 'no'} {l}.", )
    P("Mensaje al médico: «" + f["narr"].split('. ',1)[1] + "»", it=True, color=LBLUE)

# ===== 6. QUÉ PRUEBA MANDA (global) =====
H("6. ¿Qué prueba manda en cada hallazgo? (test completo)")
P("Agregando todos los pacientes de test, esto es lo que recomienda la herramienta. La radiografía lidera casi siempre "
  "(es el oráculo de su propia etiqueta), pero hay excepciones clínicamente sensatas que es donde está el valor:")
lp=D["lead_pct"]
table(["Hallazgo","% Radiografía","% ECG","% Analíticas"],
      [[ES[l],f"{lp[l]['CXR']:.0f}%",f"{lp[l]['ECG']:.0f}%",f"{lp[l]['LABS']:.0f}%"] for l in PATHO])
bullet("Edema: las analíticas lideran en el 16% de los pacientes (congestión, función renal/cardíaca en sangre).")
bullet("Cardiomegalia: el ECG lidera en el 15% (hipertrofia/eje eléctrico) — su único nicho real.")
bullet("Atelectasia, Opacidad y Derrame: domina la radiografía (>89%); las demás pruebas apenas la corrigen.")
img(FG/"exp_lead_distribucion.png",14,"Figura 5. Porcentaje de pacientes donde cada prueba lidera, por hallazgo (test).")

# ===== 7. EJEMPLOS DE PACIENTES =====
H("7. Ejemplos de pacientes (salida real de la herramienta)")
BNAME={"cxr_seguro_acierta":"Caso A — La radiografía es clara: guíate por ella",
       "defiere_labs_acierta":"Caso B — La radiografía duda y las analíticas aciertan",
       "discrepancia":"Caso C — Las pruebas discrepan: se sigue a la más fiable",
       "ecg_ignorado":"Caso D — Un ECG muy seguro pero poco hábil: se ignora",
       "multi_hallazgo":"Caso E — Paciente con varios hallazgos a la vez"}
for e in D["ejemplos"]:
    H(BNAME.get(e["bucket"],e["bucket"]),size=12,color=LBLUE,bef=10,aft=4)
    rows=[]
    for f in e["findings"]:
        rows.append([ES[f["label"]],pct(f["probs"]["CXR"]),pct(f["probs"]["ECG"]),pct(f["probs"]["LABS"]),
                     MOD_ES[f["lead"]],pct(f["comb"]),"SÍ" if f["y"]==1 else "no"])
    table(["Hallazgo","Radiog.","ECG","Analít.","Líder","Decisión","Verdad"],rows,small=True)
    foc=e.get("focus"); fsel=next((f for f in e["findings"] if f["label"]==foc),e["findings"][0])
    P("→ " + fsel["narr"], size=10, color=GRAY)
    img(FG/e["fig"],12.5)

# ===== 8. HONESTIDAD =====
H("8. Honestidad y limitaciones")
bullet("Es apoyo, NO diagnóstico: surfacea y justifica información; el médico decide. La decisión combinada acierta más "
       "que cualquier prueba sola, pero también falla (caso B incluye un hallazgo donde fiarse de las analíticas erró).")
bullet("Como las etiquetas vienen del CXR, la herramienta suele recomendar la radiografía; su valor está en los casos de "
       "CXR dudoso, sobre todo Edema/Derrame (analíticas) y Cardiomegalia (ECG).")
bullet("La fiabilidad por confianza y la habilidad se estiman en validación (n limitado, 750 pacientes): son orientación "
       "estadística, no garantías individuales.")
bullet("El test tiene 464 pacientes con etiqueta; las probabilidades son de modelos solo-CPU (CXR DenseNet121 v5, ECG "
       "ResNet1D v3.1, analíticas blend v2.2).")

# ===== 9. USO Y FUTURO =====
H("9. Cómo se usa y siguiente nivel")
P("En el notebook 05_HERRAMIENTA_DECISION, la función explicar(hadm_id) devuelve, para cualquier paciente del test, la "
  "tabla por prueba, los pesos, la decisión y el mensaje en lenguaje natural, con su figura.")
P("Mayor explicabilidad (requiere más cómputo / GPU): atribución por SHAP sobre el meta-modelo, contrafactuales («si "
  "quito la radiografía, la decisión pasa de X a Y») y atención cruzada sobre embeddings para señalar qué región de la "
  "placa, qué derivación del ECG o qué analítica concreta dispara la recomendación.")

try:
    doc.save(str(DOC)); print("Guardado:",DOC.name)
except PermissionError:
    alt=DOC.with_name("Documentación Informe Herramienta Decisión Clínica (nuevo).docx"); doc.save(str(alt)); print("BLOQUEADO. Guardado:",alt.name)
print("imgs/tablas embebidas OK")
